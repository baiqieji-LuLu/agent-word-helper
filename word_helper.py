#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""word_helper.py

General-purpose Word tool:
1. Typeset a text file or string into a DOCX;
2. Read all body text from a DOCX (including tables, in reading order);
3. Supports $...$ inline and $$...$$ display formulas (LaTeX -> OMML).

Built entirely on python-docx (>= 1.2.0) and OOXML; does not invoke Microsoft Word,
nor any COM or Office automation interfaces.

Common commands:
    python word_helper.py create --text "Hello, world" -o hello.docx
    python word_helper.py create --input notes.md -o notes.docx --title "Notes"
    python word_helper.py read report.docx
    python word_helper.py read report.docx --all --json
"""

from __future__ import annotations

import argparse
import codecs
import json
import os
import re
import sys
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

# LaTeX -> MathML -> OMML conversion dependencies; prefer the workspace .pylibs_math directory.
try:
    import latex2mathml.converter
    from lxml import etree
except ModuleNotFoundError:
    _LOCAL_MATH_LIBS = Path(__file__).resolve().parent / ".pylibs_math"
    if _LOCAL_MATH_LIBS.is_dir():
        sys.path.insert(0, str(_LOCAL_MATH_LIBS))
    try:
        import latex2mathml.converter
        from lxml import etree
    except ModuleNotFoundError:
        latex2mathml = None
        etree = None

# ---------------------------------------------------------------------------
# Adjustable layout constants
# ---------------------------------------------------------------------------

BODY_SIZE_PT = 12.0          # Body text size: 12 pt
BODY_LINE_SPACING = 1.5      # 1.5 line spacing
BODY_CJK_FONT = "SimSun"
BODY_LATIN_FONT = "Calibri"
HEADING_CJK_FONT = "Microsoft YaHei"
HEADING_LATIN_FONT = "Arial"
CODE_CJK_FONT = "DengXian"
CODE_LATIN_FONT = "Consolas"

TITLE_COLOR = RGBColor(0x1F, 0x38, 0x64)
HEADING_COLORS = {
    1: RGBColor(0x1F, 0x38, 0x64),
    2: RGBColor(0x2E, 0x54, 0x96),
    3: RGBColor(0x2E, 0x54, 0x96),
    4: RGBColor(0x40, 0x40, 0x40),
    5: RGBColor(0x40, 0x40, 0x40),
    6: RGBColor(0x40, 0x40, 0x40),
}
GRAY_TEXT = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_HEADER_FILL = "1F4E79"
TABLE_STRIPE_FILL = "F2F6FA"
INLINE_CODE_FILL = "F2F2F2"
CODE_BLOCK_FILL = "F6F8FA"

# Insert pBdr / shd elements at OOXML-required positions to avoid malformed documents.
_PBDR_SUCCESSORS = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)
_PShd_SUCCESSORS = (
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)
_RShd_SUCCESSORS = (
    "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
    "w:eastAsianLayout", "w:specVanish", "w:oMath",
)
_TCShd_SUCCESSORS = (
    "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
    "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge",
    "w:tcPrChange",
)
_TBLBORDERS_SUCCESSORS = (
    "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
    "w:tblDescription", "w:tblPrChange",
)
_TBLCELLMAR_SUCCESSORS = (
    "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange",
)
_TRCANTSplit_SUCCESSORS = (
    "w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden",
    "w:ins", "w:del", "w:trPrChange",
)
_TRHEADER_SUCCESSORS = (
    "w:tblCellSpacing", "w:jc", "w:hidden", "w:ins", "w:del",
    "w:trPrChange",
)


class WordHelperError(Exception):
    """An error whose message can be shown directly to the user."""


# Cache the XSLT transformer globally so MML2OMML.XSL is parsed only once.
_XSLT_TRANSFORMER = None


def _get_transformer():
    """Lazily load and cache the MathML -> OMML XSLT transformer."""
    global _XSLT_TRANSFORMER
    if _XSLT_TRANSFORMER is None:
        if etree is None:
            raise WordHelperError("Missing dependency: lxml; cannot convert math formulas.")
        xsl_path = Path(__file__).resolve().with_name("MML2OMML.XSL")
        if not xsl_path.is_file():
            raise WordHelperError(f"Missing formula stylesheet: {xsl_path}")
        _XSLT_TRANSFORMER = etree.XSLT(etree.parse(str(xsl_path)))
    return _XSLT_TRANSFORMER


def latex_to_omml(latex_str: str):
    """Convert LaTeX -> MathML -> OMML and return an m:oMath element."""
    if latex2mathml is None:
        raise WordHelperError(
            "Missing dependency: latex2mathml; cannot convert math formulas. "
            "Install it or place it in the workspace .pylibs_math directory."
        )
    if etree is None:
        raise WordHelperError("Missing dependency: lxml; cannot convert math formulas.")
    # Compatibility: latex2mathml does not support the aligned environment, so normalize it to the equivalent align.
    latex_str = (
        latex_str.replace(r"\begin{aligned}", r"\begin{align}")
        .replace(r"\end{aligned}", r"\end{align}")
    )
    mathml_str = latex2mathml.converter.convert(latex_str)
    mathml_tree = etree.fromstring(mathml_str.encode("utf-8"))
    omml_tree = _get_transformer()(mathml_tree)
    root = omml_tree.getroot()
    if root.tag != qn("m:oMath"):
        raise WordHelperError(f"Unexpected formula conversion result: {root.tag}")
    return root


def _math_error_detail(exc: Exception) -> str:
    """Some conversion errors have no message; fall back to the exception type name."""
    return str(exc) or type(exc).__name__


# ---------------------------------------------------------------------------
# Markdown / plain-text parsing
# ---------------------------------------------------------------------------


@dataclass
class Block:
    """Intermediate structure before layout."""

    kind: str
    text: str = ""
    level: int = 0
    number: str = ""
    lang: str = ""
    rows: list[list[str]] = field(default_factory=list)
    aligns: list[str] = field(default_factory=list)


HEADING_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
FENCE_RE = re.compile(r"^\s{0,3}(`{3,}|~{3,})\s*([^\s`]*)\s*$")
UL_RE = re.compile(r"^(\s*)[-+*]\s+(.*)$")
OL_RE = re.compile(r"^(\s*)(\d{1,9})[.)]\s+(.*)$")
QUOTE_RE = re.compile(r"^\s{0,3}>\s?(.*)$")
HR_RE = re.compile(r"^\s{0,3}(?:(?:\*\s*){3,}|(?:-\s*){3,}|(?:_\s*){3,})$")
TABLE_SEP_RE = re.compile(
    r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$"
)
INLINE_RE = re.compile(
    r"`(?P<code>[^`\n]+)`"
    r"|\[(?P<link_text>[^\]\n]+)\]\((?P<link_url>[^)\s]+)\)"
    r"|\*\*(?P<bold>[^*\n]+)\*\*"
    r"|(?<!\*)\*(?P<italic>[^*\n]+)\*(?!\*)"
    r"|~~(?P<strike>[^~\n]+)~~"
)
INLINE_HINT_RE = re.compile(
    r"\*\*[^*\n]+\*\*|`[^`\n]+`|\[[^\]\n]+\]\([^)\n]+\)"
)
# Inline math $...$, single-line display math $$...$$, and display-math fences $$.
INLINE_MATH_RE = re.compile(r"(?<!\$)\$(?!\$)(?P<math>[^$\n]+?)\$(?!\$)")
DISPLAY_MATH_RE = re.compile(r"^\s*\$\$(?P<math>.+?)\$\$\s*$")
DISPLAY_MATH_FENCE_RE = re.compile(r"^\s*\$\$\s*$")


def _normalize_newlines(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _is_cjk_or_fullwidth(ch: str) -> bool:
    code = ord(ch)
    return (
        0x3000 <= code <= 0x303F      # CJK punctuation
        or 0x3400 <= code <= 0x4DBF   # Extension A
        or 0x4E00 <= code <= 0x9FFF   # Basic block
        or 0xF900 <= code <= 0xFAFF   # Compatibility block
        or 0xFF00 <= code <= 0xFFEF   # Full-width characters
    )


def _smart_join(lines: Sequence[str]) -> str:
    """Merge Markdown soft line breaks: no space between CJK characters, otherwise one space."""
    pieces: list[str] = []
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if not pieces:
            pieces.append(line)
            continue
        prev = pieces[-1]
        if prev and (_is_cjk_or_fullwidth(prev[-1]) or _is_cjk_or_fullwidth(line[0])):
            pieces[-1] = prev + line
        else:
            pieces[-1] = prev + " " + line
    return "".join(pieces).strip()


def _starts_markdown_block(lines: Sequence[str], index: int) -> bool:
    if index >= len(lines) or not lines[index].strip():
        return True
    line = lines[index]
    if HEADING_RE.match(line) or FENCE_RE.match(line):
        return True
    if HR_RE.match(line) or QUOTE_RE.match(line):
        return True
    if DISPLAY_MATH_RE.match(line) or DISPLAY_MATH_FENCE_RE.match(line):
        return True
    if UL_RE.match(line) or OL_RE.match(line):
        return True
    if "|" in line and index + 1 < len(lines) and TABLE_SEP_RE.match(lines[index + 1]):
        return True
    return False


def looks_like_markdown(text: str) -> bool:
    """Heuristically determine whether the text looks like Markdown."""
    for line in text.splitlines():
        if (
            HEADING_RE.match(line)
            or FENCE_RE.match(line)
            or QUOTE_RE.match(line)
            or HR_RE.match(line)
            or TABLE_SEP_RE.match(line)
            or UL_RE.match(line)
            or OL_RE.match(line)
        ):
            return True
    return bool(INLINE_HINT_RE.search(text))


def resolve_format(requested: str, source_path: Path | None, text: str) -> str:
    """Resolve ``auto`` to ``markdown`` or ``plain``."""
    if requested in {"markdown", "plain"}:
        return requested
    if requested != "auto":
        raise WordHelperError(f"Unknown format: {requested}")
    if source_path and source_path.suffix.lower() in {".md", ".markdown", ".mdown"}:
        return "markdown"
    return "markdown" if looks_like_markdown(text) else "plain"


def parse_plain(text: str) -> list[Block]:
    # Plain-text mode also recognizes $$...$$ display math.
    lines = _normalize_newlines(text).split("\n")
    blocks: list[Block] = []
    index = 0
    total = len(lines)
    while index < total:
        line = lines[index]
        if not line.strip():
            index += 1
            continue

        display_match = DISPLAY_MATH_RE.match(line)
        if display_match:
            latex = display_match.group("math").strip()
            if latex:
                blocks.append(Block("math", text=latex))
            index += 1
            continue

        if DISPLAY_MATH_FENCE_RE.match(line):
            index += 1
            math_lines: list[str] = []
            while index < total and not DISPLAY_MATH_FENCE_RE.match(lines[index]):
                math_lines.append(lines[index])
                index += 1
            if index < total:
                index += 1
            latex = " ".join(part.strip() for part in math_lines if part.strip())
            if latex:
                blocks.append(Block("math", text=latex))
            continue

        blocks.append(Block("paragraph", text=line.strip()))
        index += 1
    return blocks


def _split_table_row(line: str) -> list[str]:
    value = line.strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|") and not value.endswith("\\|"):
        value = value[:-1]
    cells = re.split(r"(?<!\\)\|", value)
    return [cell.strip().replace("\\|", "|") for cell in cells]


def _parse_table(lines: Sequence[str], start: int) -> tuple[Block, int]:
    header = _split_table_row(lines[start])
    separator = _split_table_row(lines[start + 1])
    ncols = max(len(header), len(separator), 1)
    header = header + [""] * (ncols - len(header))

    aligns: list[str] = []
    for index in range(ncols):
        spec = separator[index].strip() if index < len(separator) else ""
        if spec.startswith(":") and spec.endswith(":"):
            aligns.append("center")
        elif spec.endswith(":"):
            aligns.append("right")
        else:
            aligns.append("left")

    rows: list[list[str]] = [header]
    index = start + 2
    while index < len(lines):
        line = lines[index]
        if not line.strip() or "|" not in line:
            break
        cells = _split_table_row(line)
        if len(cells) > ncols:
            extra = len(cells) - ncols
            for row in rows:
                row.extend([""] * extra)
            aligns.extend(["left"] * extra)
            ncols = len(cells)
        elif len(cells) < ncols:
            cells.extend([""] * (ncols - len(cells)))
        rows.append(cells)
        index += 1

    return Block("table", rows=rows, aligns=aligns), index


def parse_markdown(text: str) -> list[Block]:
    lines = _normalize_newlines(text).split("\n")
    blocks: list[Block] = []
    index = 0
    total = len(lines)

    while index < total:
        line = lines[index]

        if not line.strip():
            index += 1
            continue

        # Single-line display math $$...$$.
        display_match = DISPLAY_MATH_RE.match(line)
        if display_match:
            latex = display_match.group("math").strip()
            if latex:
                blocks.append(Block("math", text=latex))
            index += 1
            continue

        # Fenced display math (a line containing only $$ wraps multi-line LaTeX).
        if DISPLAY_MATH_FENCE_RE.match(line):
            index += 1
            math_lines: list[str] = []
            while index < total and not DISPLAY_MATH_FENCE_RE.match(lines[index]):
                math_lines.append(lines[index])
                index += 1
            if index < total:
                index += 1
            latex = " ".join(part.strip() for part in math_lines if part.strip())
            if latex:
                blocks.append(Block("math", text=latex))
            continue

        fence_match = FENCE_RE.match(line)
        if fence_match:
            fence = fence_match.group(1)
            fence_char = fence[0]
            fence_length = len(fence)
            lang = fence_match.group(2) or ""
            index += 1
            code_lines: list[str] = []
            while index < total:
                stripped = lines[index].strip()
                if (
                    len(stripped) >= fence_length
                    and stripped
                    and all(ch == fence_char for ch in stripped)
                ):
                    index += 1
                    break
                code_lines.append(lines[index])
                index += 1
            blocks.append(Block("code", text="\n".join(code_lines), lang=lang))
            continue

        if (
            "|" in line
            and index + 1 < total
            and TABLE_SEP_RE.match(lines[index + 1])
        ):
            table_block, index = _parse_table(lines, index)
            blocks.append(table_block)
            continue

        heading_match = HEADING_RE.match(line)
        if heading_match:
            heading_text = heading_match.group(2).strip()
            if heading_text:
                blocks.append(
                    Block(
                        "heading",
                        text=heading_text,
                        level=len(heading_match.group(1)),
                    )
                )
            index += 1
            continue

        if HR_RE.match(line):
            blocks.append(Block("hr"))
            index += 1
            continue

        if QUOTE_RE.match(line):
            quote_lines: list[str] = []
            while index < total:
                quote_match = QUOTE_RE.match(lines[index])
                if not quote_match:
                    break
                quote_lines.append(quote_match.group(1))
                index += 1
            quote_text = _smart_join(quote_lines)
            if quote_text:
                blocks.append(Block("quote", text=quote_text))
            continue

        bullet_match = UL_RE.match(line)
        if bullet_match:
            indent = len(bullet_match.group(1).replace("\t", "    "))
            level = min(indent // 2, 3)
            item_text = bullet_match.group(2).strip()
            if item_text:
                blocks.append(
                    Block("bullet", text=item_text, level=level)
                )
            index += 1
            continue

        number_match = OL_RE.match(line)
        if number_match:
            indent = len(number_match.group(1).replace("\t", "    "))
            level = min(indent // 2, 3)
            item_text = number_match.group(3).strip()
            if item_text:
                blocks.append(
                    Block(
                        "number",
                        text=item_text,
                        level=level,
                        number=number_match.group(2),
                    )
                )
            index += 1
            continue

        paragraph_lines: list[str] = []
        while index < total:
            if not lines[index].strip():
                break
            if index > 0 and paragraph_lines and _starts_markdown_block(lines, index):
                break
            paragraph_lines.append(lines[index])
            index += 1
        paragraph_text = _smart_join(paragraph_lines)
        if paragraph_text:
            blocks.append(Block("paragraph", text=paragraph_text))

    return blocks


def parse_text(text: str, fmt: str) -> list[Block]:
    if fmt == "markdown":
        return parse_markdown(text)
    return parse_plain(text)


# ---------------------------------------------------------------------------
# OOXML helpers
# ---------------------------------------------------------------------------


def _set_style_font(
    style,
    *,
    latin: str,
    cjk: str,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: RGBColor | None = None,
) -> None:
    style.font.name = latin
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), cjk)
    r_fonts.set(qn("w:cs"), latin)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color


def _set_run_font(
    run,
    *,
    latin: str,
    cjk: str,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: RGBColor | None = None,
) -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), cjk)
    r_fonts.set(qn("w:cs"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _set_pbdr(p_pr, *, edge: str, color: str, size: int, space: int) -> None:
    p_bdr = p_pr.find(qn("w:pBdr"))
    is_new = p_bdr is None
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")

    existing_edge = p_bdr.find(qn(f"w:{edge}"))
    if existing_edge is not None:
        p_bdr.remove(existing_edge)
    border = OxmlElement(f"w:{edge}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)
    p_bdr.append(border)

    if is_new:
        p_pr.insert_element_before(p_bdr, *_PBDR_SUCCESSORS)


def _set_p_shading(p_pr, fill: str) -> None:
    existing = p_pr.find(qn("w:shd"))
    if existing is not None:
        p_pr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.insert_element_before(shd, *_PShd_SUCCESSORS)


def _set_run_shading(run, fill: str) -> None:
    r_pr = run._element.get_or_add_rPr()
    existing = r_pr.find(qn("w:shd"))
    if existing is not None:
        r_pr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    r_pr.insert_element_before(shd, *_RShd_SUCCESSORS)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:shd"))
    if existing is not None:
        tc_pr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.insert_element_before(shd, *_TCShd_SUCCESSORS)


def _set_table_borders(table, color: str = "BFBFBF") -> None:
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn("w:tblBorders"))
    if existing is not None:
        tbl_pr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tbl_pr.insert_element_before(borders, *_TBLBORDERS_SUCCESSORS)


def _set_table_cell_margins(table) -> None:
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn("w:tblCellMar"))
    if existing is not None:
        tbl_pr.remove(existing)
    margins = OxmlElement("w:tblCellMar")
    for edge, width in (("top", 50), ("left", 120), ("bottom", 50), ("right", 120)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(width))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tbl_pr.insert_element_before(margins, *_TBLCELLMAR_SUCCESSORS)


def _mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.insert_element_before(header, *_TRHEADER_SUCCESSORS)


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.insert_element_before(cant_split, *_TRCANTSplit_SUCCESSORS)


def _set_body_first_line_indent(paragraph, chars: float = 2.0) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    ind = p_pr.get_or_add_ind()
    ind.set(qn("w:firstLineChars"), str(int(chars * 100)))
    ind.set(qn("w:firstLine"), str(int(round(BODY_SIZE_PT * chars * 20))))
    ind.set(qn("w:left"), ind.get(qn("w:left"), "0"))
    ind.set(qn("w:right"), ind.get(qn("w:right"), "0"))


def _add_field(
    paragraph: Paragraph,
    instruction: str,
    placeholder: str,
    *,
    size: float,
    color: RGBColor,
) -> None:
    def field_run(field_type: str):
        run = paragraph.add_run()
        field_char = OxmlElement("w:fldChar")
        field_char.set(qn("w:fldCharType"), field_type)
        run._r.append(field_char)
        return run

    field_run("begin")

    instruction_run = paragraph.add_run()
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = f" {instruction} "
    instruction_run._r.append(instr_text)

    field_run("separate")

    result_run = paragraph.add_run(placeholder)
    _set_run_font(
        result_run,
        latin=BODY_LATIN_FONT,
        cjk=BODY_CJK_FONT,
        size=size,
        color=color,
    )

    field_run("end")


# ---------------------------------------------------------------------------
# DOCX generation
# ---------------------------------------------------------------------------


def _configure_section(doc) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def _configure_styles(doc) -> None:
    normal = doc.styles["Normal"]
    _set_style_font(
        normal,
        latin=BODY_LATIN_FONT,
        cjk=BODY_CJK_FONT,
        size=BODY_SIZE_PT,
    )
    normal.paragraph_format.line_spacing = BODY_LINE_SPACING
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    _set_style_font(
        title,
        latin=HEADING_LATIN_FONT,
        cjk=HEADING_CJK_FONT,
        size=24,
        bold=True,
        color=TITLE_COLOR,
    )
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.line_spacing = 1.15
    # The default template Title style has a bottom border; the title block
    # draws its line after the subtitle, so remove the template border here to avoid a double line.
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = doc.styles["Subtitle"]
    _set_style_font(
        subtitle,
        latin=BODY_LATIN_FONT,
        cjk=BODY_CJK_FONT,
        size=10.5,
        bold=False,
        italic=False,
        color=GRAY_TEXT,
    )
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(2)
    subtitle.paragraph_format.line_spacing = 1.0

    heading_sizes = {1: 18, 2: 15, 3: 13, 4: 12, 5: 11, 6: 11}
    before = {1: 18, 2: 14, 3: 12, 4: 10, 5: 8, 6: 8}
    after = {1: 8, 2: 6, 3: 4, 4: 3, 5: 2, 6: 2}
    for level in range(1, 7):
        style = doc.styles[f"Heading {level}"]
        _set_style_font(
            style,
            latin=HEADING_LATIN_FONT,
            cjk=HEADING_CJK_FONT,
            size=heading_sizes[level],
            bold=True,
            color=HEADING_COLORS[level],
        )
        style.paragraph_format.space_before = Pt(before[level])
        style.paragraph_format.space_after = Pt(after[level])
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Pt(0)
        if level == 1:
            _set_pbdr(
                style.element.get_or_add_pPr(),
                edge="bottom",
                color="8EAADB",
                size=8,
                space=4,
            )

    for style_name in ("List Bullet", "List Bullet 2", "List Bullet 3"):
        style = doc.styles[style_name]
        _set_style_font(
            style,
            latin=BODY_LATIN_FONT,
            cjk=BODY_CJK_FONT,
            size=BODY_SIZE_PT,
        )
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.4
        style.paragraph_format.first_line_indent = Pt(0)

    quote = doc.styles["Quote"]
    _set_style_font(
        quote,
        latin=BODY_LATIN_FONT,
        cjk=BODY_CJK_FONT,
        size=11,
        bold=False,
        italic=False,
        color=GRAY_TEXT,
    )
    quote.paragraph_format.left_indent = Cm(0.75)
    quote.paragraph_format.right_indent = Cm(0.3)
    quote.paragraph_format.first_line_indent = Pt(0)
    quote.paragraph_format.space_before = Pt(4)
    quote.paragraph_format.space_after = Pt(6)
    quote.paragraph_format.line_spacing = 1.4
    _set_pbdr(
        quote.element.get_or_add_pPr(),
        edge="left",
        color="8EAADB",
        size=18,
        space=8,
    )

    try:
        code_style = doc.styles["Code Block"]
    except KeyError:
        code_style = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code_style.base_style = doc.styles["Normal"]
        code_style.quick_style = True
    _set_style_font(
        code_style,
        latin=CODE_LATIN_FONT,
        cjk=CODE_CJK_FONT,
        size=10,
        color=RGBColor(0x24, 0x29, 0x2F),
    )
    code_style.paragraph_format.left_indent = Cm(0.3)
    code_style.paragraph_format.right_indent = Cm(0.3)
    code_style.paragraph_format.first_line_indent = Pt(0)
    code_style.paragraph_format.space_before = Pt(6)
    code_style.paragraph_format.space_after = Pt(6)
    code_style.paragraph_format.line_spacing = 1.0
    code_ppr = code_style.element.get_or_add_pPr()
    _set_p_shading(code_ppr, CODE_BLOCK_FILL)
    _set_pbdr(code_ppr, edge="top", color="D0D7DE", size=4, space=4)
    _set_pbdr(code_ppr, edge="left", color="D0D7DE", size=4, space=4)


def _setup_header_footer(doc, display_title: str) -> None:
    section = doc.sections[0]

    header_paragraph = section.header.paragraphs[0]
    header_paragraph.text = ""
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_paragraph.paragraph_format.space_before = Pt(0)
    header_paragraph.paragraph_format.space_after = Pt(2)
    _add_header_text(header_paragraph, _shorten(display_title, 70))
    _set_pbdr(
        header_paragraph._p.get_or_add_pPr(),
        edge="bottom",
        color="D9D9D9",
        size=4,
        space=3,
    )

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.text = ""
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.paragraph_format.space_before = Pt(2)
    footer_paragraph.paragraph_format.space_after = Pt(0)
    left_dash = footer_paragraph.add_run("— ")
    _set_run_font(
        left_dash,
        latin=BODY_LATIN_FONT,
        cjk=BODY_CJK_FONT,
        size=9,
        color=RGBColor(0x80, 0x80, 0x80),
    )
    _add_field(
        footer_paragraph,
        "PAGE",
        "1",
        size=9,
        color=RGBColor(0x80, 0x80, 0x80),
    )
    right_dash = footer_paragraph.add_run(" —")
    _set_run_font(
        right_dash,
        latin=BODY_LATIN_FONT,
        cjk=BODY_CJK_FONT,
        size=9,
        color=RGBColor(0x80, 0x80, 0x80),
    )


def _shorten(text: str, limit: int) -> str:
    compact = " ".join(text.split())
    if len(compact) <= limit:
        return compact
    return compact[: limit - 1] + "…"


def _set_omml_size(omml_root, size_pt: float) -> None:
    """Set the font size of runs inside OMML (e.g., for small header text)."""
    value = str(int(round(size_pt * 2)))
    for run in omml_root.iter(qn("m:r")):
        r_pr = run.find(qn("w:rPr"))
        if r_pr is None:
            r_pr = OxmlElement("w:rPr")
            insert_at = 1 if len(run) and run[0].tag == qn("m:rPr") else 0
            run.insert(insert_at, r_pr)
        for tag in ("w:sz", "w:szCs"):
            element = r_pr.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                r_pr.append(element)
            element.set(qn("w:val"), value)


def _add_header_text(paragraph: Paragraph, text: str) -> None:
    """Header text, supporting inline $...$ math."""
    position = 0
    for match in INLINE_MATH_RE.finditer(text):
        if match.start() > position:
            _add_text_run(
                paragraph,
                text[position : match.start()],
                size=9,
                color=RGBColor(0x80, 0x80, 0x80),
                latin_font=BODY_LATIN_FONT,
                cjk_font=HEADING_CJK_FONT,
            )
        try:
            omml = latex_to_omml(match.group("math"))
        except WordHelperError:
            raise
        except Exception as exc:
            raise WordHelperError(
                f"Formula conversion failed: {match.group('math')} ({_math_error_detail(exc)})"
            ) from exc
        _set_omml_size(omml, 9)
        paragraph._p.append(omml)
        position = match.end()

    if position < len(text):
        _add_text_run(
            paragraph,
            text[position:],
            size=9,
            color=RGBColor(0x80, 0x80, 0x80),
            latin_font=BODY_LATIN_FONT,
            cjk_font=HEADING_CJK_FONT,
        )


def _add_text_run(
    paragraph: Paragraph,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor | None = None,
    size: float | None = None,
    latin_font: str = BODY_LATIN_FONT,
    cjk_font: str = BODY_CJK_FONT,
    strike: bool = False,
) -> None:
    chunks = text.split("\n")
    for index, chunk in enumerate(chunks):
        if index:
            paragraph.add_run().add_break()
        if chunk:
            run = paragraph.add_run(chunk)
            _set_run_font(
                run,
                latin=latin_font,
                cjk=cjk_font,
                size=size,
                bold=bold,
                italic=italic,
                color=color,
            )
            if strike:
                run.font.strike = True


def _add_hyperlink(
    paragraph: Paragraph,
    url: str,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    size: float | None = None,
    latin_font: str = BODY_LATIN_FONT,
    cjk_font: str = BODY_CJK_FONT,
) -> None:
    run = paragraph.add_run(text)
    _set_run_font(
        run,
        latin=latin_font,
        cjk=cjk_font,
        size=size,
        bold=bold,
        italic=italic,
        color=RGBColor(0x05, 0x63, 0xC1),
    )
    run.font.underline = True

    relationship_id = paragraph.part.relate_to(
        url, RT.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    hyperlink.set(qn("w:history"), "1")
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def _add_inline(
    paragraph: Paragraph,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor | None = None,
    size: float | None = None,
    latin_font: str = BODY_LATIN_FONT,
    cjk_font: str = BODY_CJK_FONT,
) -> None:
    """Split $...$ math first, then render the remaining text as inline Markdown."""
    position = 0
    for match in INLINE_MATH_RE.finditer(text):
        if match.start() > position:
            _add_inline_markdown(
                paragraph,
                text[position : match.start()],
                bold=bold,
                italic=italic,
                color=color,
                size=size,
                latin_font=latin_font,
                cjk_font=cjk_font,
            )
        try:
            paragraph._p.append(latex_to_omml(match.group("math")))
        except WordHelperError:
            raise
        except Exception as exc:
            raise WordHelperError(
                f"Formula conversion failed: {match.group('math')} ({_math_error_detail(exc)})"
            ) from exc
        position = match.end()

    if position < len(text):
        _add_inline_markdown(
            paragraph,
            text[position:],
            bold=bold,
            italic=italic,
            color=color,
            size=size,
            latin_font=latin_font,
            cjk_font=cjk_font,
        )


def _add_inline_markdown(
    paragraph: Paragraph,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor | None = None,
    size: float | None = None,
    latin_font: str = BODY_LATIN_FONT,
    cjk_font: str = BODY_CJK_FONT,
) -> None:
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            _add_text_run(
                paragraph,
                text[position : match.start()],
                bold=bold,
                italic=italic,
                color=color,
                size=size,
                latin_font=latin_font,
                cjk_font=cjk_font,
            )

        if match.group("code") is not None:
            code_run = paragraph.add_run(match.group("code"))
            _set_run_font(
                code_run,
                latin=CODE_LATIN_FONT,
                cjk=CODE_CJK_FONT,
                size=(size - 0.5) if size else BODY_SIZE_PT - 0.5,
                color=RGBColor(0x9C, 0x27, 0x2E),
            )
            _set_run_shading(code_run, INLINE_CODE_FILL)
        elif match.group("link_text") is not None:
            _add_hyperlink(
                paragraph,
                match.group("link_url"),
                match.group("link_text"),
                bold=bold,
                italic=italic,
                size=size,
                latin_font=latin_font,
                cjk_font=cjk_font,
            )
        elif match.group("bold") is not None:
            _add_text_run(
                paragraph,
                match.group("bold"),
                bold=True,
                italic=italic,
                color=color,
                size=size,
                latin_font=latin_font,
                cjk_font=cjk_font,
            )
        elif match.group("italic") is not None:
            _add_text_run(
                paragraph,
                match.group("italic"),
                bold=bold,
                italic=True,
                color=color,
                size=size,
                latin_font=latin_font,
                cjk_font=cjk_font,
            )
        elif match.group("strike") is not None:
            _add_text_run(
                paragraph,
                match.group("strike"),
                bold=bold,
                italic=italic,
                color=color,
                size=size,
                latin_font=latin_font,
                cjk_font=cjk_font,
                strike=True,
            )

        position = match.end()

    if position < len(text):
        _add_text_run(
            paragraph,
            text[position:],
            bold=bold,
            italic=italic,
            color=color,
            size=size,
            latin_font=latin_font,
            cjk_font=cjk_font,
        )


def _add_text_with_math(paragraph: Paragraph, text: str) -> None:
    """Render inline math in plain-text mode without interpreting Markdown."""
    position = 0
    for match in INLINE_MATH_RE.finditer(text):
        if match.start() > position:
            _add_text_run(paragraph, text[position : match.start()])
        try:
            paragraph._p.append(latex_to_omml(match.group("math")))
        except WordHelperError:
            raise
        except Exception as exc:
            raise WordHelperError(
                f"Formula conversion failed: {match.group('math')} ({_math_error_detail(exc)})"
            ) from exc
        position = match.end()

    if position < len(text):
        _add_text_run(paragraph, text[position:])


def _add_title_block(doc, title: str, subtitle: str | None) -> None:
    title_paragraph = doc.add_paragraph(style="Title")
    _add_text_run(
        title_paragraph,
        title,
        bold=True,
        size=24,
        color=TITLE_COLOR,
        latin_font=HEADING_LATIN_FONT,
        cjk_font=HEADING_CJK_FONT,
    )
    if subtitle:
        subtitle_paragraph = doc.add_paragraph(style="Subtitle")
        _add_text_run(
            subtitle_paragraph,
            subtitle,
            color=GRAY_TEXT,
            size=10.5,
            latin_font=BODY_LATIN_FONT,
            cjk_font=BODY_CJK_FONT,
        )
    _add_horizontal_rule(doc, color="8EAADB", space_before=2, space_after=14)


def _add_horizontal_rule(
    doc,
    *,
    color: str = "D9D9D9",
    space_before: int = 2,
    space_after: int = 8,
) -> Paragraph:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = Pt(1)
    _set_pbdr(
        paragraph._p.get_or_add_pPr(),
        edge="bottom",
        color=color,
        size=6,
        space=1,
    )
    return paragraph


def _add_heading(doc, text: str, level: int) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {min(max(level, 1), 6)}")
    _add_inline(
        paragraph,
        text,
        latin_font=HEADING_LATIN_FONT,
        cjk_font=HEADING_CJK_FONT,
    )


def _add_body_paragraph(doc, text: str, *, use_inline: bool = True) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_body_first_line_indent(paragraph)
    if use_inline:
        _add_inline(paragraph, text)
    else:
        _add_text_with_math(paragraph, text)


def _add_bullet(doc, text: str, level: int) -> None:
    style_name = "List Bullet" if level == 0 else f"List Bullet {min(level + 1, 3)}"
    paragraph = doc.add_paragraph(style=style_name)
    paragraph.paragraph_format.space_after = Pt(2)
    if level > 2:
        paragraph.paragraph_format.left_indent = Cm(0.75 * (level + 1))
    _add_inline(paragraph, text)


def _add_numbered(doc, text: str, level: int, number: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    left = 0.75 + 0.75 * level
    paragraph.paragraph_format.left_indent = Cm(left)
    paragraph.paragraph_format.first_line_indent = Cm(-0.55)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.4
    number_run = paragraph.add_run(f"{number}. ")
    _set_run_font(
        number_run,
        latin=BODY_LATIN_FONT,
        cjk=BODY_CJK_FONT,
        size=BODY_SIZE_PT,
    )
    _add_inline(paragraph, text)


def _add_quote(doc, text: str) -> None:
    paragraph = doc.add_paragraph(style="Quote")
    _add_inline(paragraph, text, color=GRAY_TEXT)


def _add_code_block(doc, text: str) -> None:
    paragraph = doc.add_paragraph(style="Code Block")
    lines = text.split("\n") if text else [""]
    for index, line in enumerate(lines):
        if line:
            run = paragraph.add_run(line)
            _set_run_font(
                run,
                latin=CODE_LATIN_FONT,
                cjk=CODE_CJK_FONT,
                size=10,
                color=RGBColor(0x24, 0x29, 0x2F),
            )
        if index < len(lines) - 1:
            paragraph.add_run().add_break()


def _alignment_for_cell(value: str) -> WD_ALIGN_PARAGRAPH:
    if value == "center":
        return WD_ALIGN_PARAGRAPH.CENTER
    if value == "right":
        return WD_ALIGN_PARAGRAPH.RIGHT
    return WD_ALIGN_PARAGRAPH.LEFT


def _add_table(doc, block: Block) -> None:
    if not block.rows:
        return
    ncols = max((len(row) for row in block.rows), default=0)
    if ncols == 0:
        return

    table = doc.add_table(rows=len(block.rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_table_borders(table)
    _set_table_cell_margins(table)

    for row_index, row in enumerate(block.rows):
        for column_index in range(ncols):
            cell_text = row[column_index] if column_index < len(row) else ""
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_paragraph = cell.paragraphs[0]
            cell_paragraph.paragraph_format.space_before = Pt(2)
            cell_paragraph.paragraph_format.space_after = Pt(2)
            cell_paragraph.paragraph_format.line_spacing = 1.15
            cell_paragraph.paragraph_format.first_line_indent = Pt(0)
            align = (
                block.aligns[column_index]
                if column_index < len(block.aligns)
                else "left"
            )
            cell_paragraph.alignment = _alignment_for_cell(align)

            if row_index == 0:
                _set_cell_shading(cell, TABLE_HEADER_FILL)
                _add_inline(
                    cell_paragraph,
                    cell_text,
                    bold=True,
                    color=WHITE,
                    size=10.5,
                )
            else:
                if row_index % 2 == 0:
                    _set_cell_shading(cell, TABLE_STRIPE_FILL)
                _add_inline(cell_paragraph, cell_text, size=10.5)

    _mark_header_row(table.rows[0])
    for row in table.rows:
        _prevent_row_split(row)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(2)


def _add_display_math(doc, latex: str) -> None:
    """Render $$...$$ as a centered display Word equation."""
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    try:
        omml = latex_to_omml(latex)
    except WordHelperError:
        raise
    except Exception as exc:
        raise WordHelperError(
            f"Formula conversion failed: {latex} ({_math_error_detail(exc)})"
        ) from exc
    math_paragraph = OxmlElement("m:oMathPara")
    math_paragraph.append(omml)
    paragraph._p.append(math_paragraph)


def _add_blocks(
    doc, blocks: Sequence[Block], *, use_inline: bool = True
) -> None:
    for block in blocks:
        if block.kind == "heading":
            _add_heading(doc, block.text, block.level)
        elif block.kind == "paragraph":
            _add_body_paragraph(doc, block.text, use_inline=use_inline)
        elif block.kind == "bullet":
            _add_bullet(doc, block.text, block.level)
        elif block.kind == "number":
            _add_numbered(doc, block.text, block.level, block.number)
        elif block.kind == "quote":
            _add_quote(doc, block.text)
        elif block.kind == "code":
            _add_code_block(doc, block.text)
        elif block.kind == "table":
            _add_table(doc, block)
        elif block.kind == "math":
            _add_display_math(doc, block.text)
        elif block.kind == "hr":
            _add_horizontal_rule(doc)


def _set_core_properties(doc, title: str, author: str | None) -> None:
    properties = doc.core_properties
    properties.title = title
    properties.author = author or ""
    properties.last_modified_by = "word_helper.py"
    properties.subject = "Document generated by word_helper.py"
    properties.comments = "Generated with python-docx; no COM interfaces used."
    properties.language = "zh-CN"


def _build_document(
    text: str,
    *,
    fmt: str,
    title: str | None,
    author: str | None,
    date_text: str | None,
    source_path: Path | None,
):
    blocks = parse_text(text, fmt)
    if not blocks:
        raise WordHelperError("No non-empty text to write.")

    explicit_title = title.strip() if title else ""
    first_heading = next(
        (
            block.text.strip()
            for block in blocks
            if block.kind == "heading" and block.text.strip()
        ),
        "",
    )
    fallback_title = source_path.stem if source_path else "Untitled Document"
    display_title = explicit_title or first_heading or fallback_title

    doc = Document()
    _configure_section(doc)
    _configure_styles(doc)
    _setup_header_footer(doc, display_title)

    if explicit_title:
        if (
            blocks
            and blocks[0].kind == "heading"
            and blocks[0].text.strip() == explicit_title
        ):
            blocks = blocks[1:]
        subtitle_parts = [part.strip() for part in (author, date_text) if part and part.strip()]
        _add_title_block(doc, explicit_title, " · ".join(subtitle_parts) or None)

    _add_blocks(doc, blocks, use_inline=(fmt != "plain"))
    _set_core_properties(doc, display_title, author)
    return doc, display_title


def _save_docx(doc, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=".word_helper_",
        suffix=".docx",
        dir=str(output.parent),
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    try:
        doc.save(str(temporary_path))
        # Reopen once to confirm the ZIP, relationships, and XML can be read by python-docx.
        Document(str(temporary_path))
        os.replace(str(temporary_path), str(output))
    except Exception as exc:
        raise WordHelperError(f"Failed to save DOCX: {exc}") from exc
    finally:
        if temporary_path.exists():
            try:
                temporary_path.unlink()
            except OSError:
                pass


def create_docx(
    text: str,
    output: str | Path,
    *,
    title: str | None = None,
    author: str | None = None,
    date_text: str | None = None,
    fmt: str = "auto",
    source_path: str | Path | None = None,
    overwrite: bool = True,
) -> Path:
    """Typeset a string into DOCX and return the output path."""
    if not isinstance(text, str):
        raise WordHelperError("text must be a string.")
    normalized = _normalize_newlines(text)
    if not normalized.strip():
        raise WordHelperError("Input text is empty.")

    output_path = Path(output).expanduser().resolve()
    if output_path.is_dir():
        raise WordHelperError(f"Output path is a directory: {output_path}")
    if output_path.suffix.lower() != ".docx":
        raise WordHelperError("Output file must have a .docx extension.")
    if output_path.exists() and not overwrite:
        raise WordHelperError(f"Output file already exists: {output_path}")

    source = Path(source_path).expanduser().resolve() if source_path else None
    resolved_format = resolve_format(fmt, source, normalized)
    doc, _ = _build_document(
        normalized,
        fmt=resolved_format,
        title=title,
        author=author,
        date_text=date_text,
        source_path=source,
    )
    _save_docx(doc, output_path)
    return output_path


# ---------------------------------------------------------------------------
# DOCX text extraction
# ---------------------------------------------------------------------------


def _paragraph_text(paragraph_element) -> str:
    """Extract paragraph text in document order, including hyperlinks, field results, and text boxes."""
    parts: list[str] = []

    def walk(element) -> None:
        for child in element:
            tag = child.tag
            if tag == qn("w:t"):
                parts.append(child.text or "")
            elif tag == qn("w:tab"):
                parent = child.getparent()
                if parent is not None and parent.tag == qn("w:r"):
                    parts.append("\t")
            elif tag in (qn("w:br"), qn("w:cr")):
                parts.append("\n")
            elif tag == qn("w:p"):
                if parts and not parts[-1].endswith("\n"):
                    parts.append("\n")
                walk(child)
                if parts and not parts[-1].endswith("\n"):
                    parts.append("\n")
            else:
                walk(child)

    walk(paragraph_element)
    return "".join(parts)


def _table_rows(table: Table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:
            identity = id(cell._tc)
            if identity in seen:
                continue
            seen.add(identity)
            cells.append(_cell_text(cell))
        rows.append(cells)
    return rows


def _cell_text(cell) -> str:
    parts: list[str] = []
    for block in _container_blocks(cell):
        if block["type"] == "paragraph":
            parts.append(block["text"])
        else:
            parts.append(_render_table_rows(block["rows"]))
    return "\n".join(part for part in parts if part.strip())


def _container_blocks(container) -> list[dict]:
    blocks: list[dict] = []
    for item in container.iter_inner_content():
        if isinstance(item, Paragraph):
            text = _paragraph_text(item._p)
            if text.strip():
                blocks.append({"type": "paragraph", "text": text})
        elif isinstance(item, Table):
            rows = _table_rows(item)
            if any(cell.strip() for row in rows for cell in row):
                blocks.append({"type": "table", "rows": rows})
    return blocks


def _render_table_rows(rows: Sequence[Sequence[str]]) -> str:
    lines = []
    for row in rows:
        lines.append(
            " | ".join(cell.replace("\n", " / ") for cell in row)
        )
    return "\n".join(lines)


def _render_blocks(blocks: Sequence[dict]) -> str:
    chunks: list[str] = []
    for block in blocks:
        if block["type"] == "paragraph":
            chunks.append(block["text"])
        elif block["type"] == "table":
            chunks.append(_render_table_rows(block["rows"]))
    return "\n\n".join(chunks)


def _collect_peripheral(doc, kind: str) -> list[dict]:
    entries: list[dict] = []
    seen: set[str] = set()
    multiple_sections = len(doc.sections) > 1
    odd_even = bool(doc.settings.odd_and_even_pages_header_footer)

    for section_index, section in enumerate(doc.sections, start=1):
        suffix = f" (section {section_index})" if multiple_sections else ""
        if kind == "header":
            candidates = [
                (f"Header{suffix}", section.header),
                (f"First-page header{suffix}", section.first_page_header),
                (f"Even-page header{suffix}", section.even_page_header),
            ]
        else:
            candidates = [
                (f"Footer{suffix}", section.footer),
                (f"First-page footer{suffix}", section.first_page_footer),
                (f"Even-page footer{suffix}", section.even_page_footer),
            ]

        for label, part in candidates:
            if label.startswith("First-page") and not section.different_first_page_header_footer:
                continue
            if label.startswith("Even-page") and not odd_even:
                continue
            blocks = _container_blocks(part)
            text = _render_blocks(blocks)
            signature = text.strip()
            if not signature or signature in seen:
                continue
            seen.add(signature)
            entries.append(
                {
                    "type": kind,
                    "label": label,
                    "section": section_index,
                    "blocks": blocks,
                    "text": text,
                }
            )
    return entries


def _collect_notes(docx_path: Path, part_name: str, label: str) -> list[dict]:
    """Read footnote and endnote text from footnotes.xml / endnotes.xml."""
    try:
        with zipfile.ZipFile(docx_path) as package:
            try:
                xml_bytes = package.read(part_name)
            except KeyError:
                return []
    except (zipfile.BadZipFile, OSError):
        return []

    try:
        root = parse_xml(xml_bytes)
    except Exception:
        return []

    entries: list[dict] = []
    for note in root:
        if note.tag not in (qn("w:footnote"), qn("w:endnote")):
            continue
        note_id = note.get(qn("w:id"))
        if note_id in {"-1", "0"}:
            continue
        text = _paragraph_text(note).strip()
        if not text:
            continue
        block = {"type": "paragraph", "text": text}
        entries.append(
            {
                "type": "footnote" if note.tag == qn("w:footnote") else "endnote",
                "label": f"{label} {note_id}",
                "id": note_id,
                "blocks": [block],
                "text": text,
            }
        )
    return entries


def extract_docx(
    path: str | Path,
    *,
    include_all_parts: bool = False,
) -> dict:
    """Read a DOCX and return body, tables, and optionally headers, footers, footnotes, and endnotes.

    By default only the body and tables are read, in document order, and stored in ``body``.
    """
    docx_path = Path(path).expanduser().resolve()
    if not docx_path.is_file():
        raise WordHelperError(f"File not found: {docx_path}")
    if docx_path.suffix.lower() != ".docx":
        raise WordHelperError("Only .docx files can be read.")
    try:
        doc = Document(str(docx_path))
    except Exception as exc:
        raise WordHelperError(f"Unable to open DOCX: {exc}") from exc

    body = _container_blocks(doc)
    body_text = _render_blocks(body)
    result: dict = {
        "path": str(docx_path),
        "body": body,
        "body_text": body_text,
        "text": body_text,
    }

    if include_all_parts:
        headers = _collect_peripheral(doc, "header")
        footers = _collect_peripheral(doc, "footer")
        footnotes = _collect_notes(docx_path, "word/footnotes.xml", "Footnote")
        endnotes = _collect_notes(docx_path, "word/endnotes.xml", "Endnote")
        result["headers"] = headers
        result["footers"] = footers
        result["footnotes"] = footnotes
        result["endnotes"] = endnotes

        chunks: list[str] = []
        if body_text:
            chunks.append(body_text)
        for entries in (headers, footers, footnotes, endnotes):
            for entry in entries:
                chunks.append(f"--- {entry['label']} ---\n{entry['text']}")
        result["text"] = "\n\n".join(chunks)

    return result


def read_docx_text(
    path: str | Path,
    *,
    include_all_parts: bool = False,
) -> str:
    """Read the text content of a DOCX. Only the body and tables are read by default."""
    return extract_docx(
        path, include_all_parts=include_all_parts
    )["text"]


# ---------------------------------------------------------------------------
# Input / CLI
# ---------------------------------------------------------------------------


def _read_text_file(path: Path, encoding: str | None) -> str:
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise WordHelperError(f"Unable to read text file: {exc}") from exc

    candidates: list[str]
    if encoding:
        candidates = [encoding]
    elif data.startswith(codecs.BOM_UTF8):
        candidates = ["utf-8-sig"]
    elif data.startswith(codecs.BOM_UTF16_LE) or data.startswith(
        codecs.BOM_UTF16_BE
    ):
        candidates = ["utf-16"]
    else:
        candidates = ["utf-8", "gb18030"]

    last_error: UnicodeDecodeError | None = None
    for candidate in candidates:
        try:
            return _normalize_newlines(data.decode(candidate))
        except UnicodeDecodeError as exc:
            last_error = exc
    raise WordHelperError(
        "Unable to decode text file; specify an encoding with --encoding (e.g. utf-8 or gbk)."
        + (f" Last error: {last_error}" if last_error else "")
    )


def _read_input(args) -> tuple[str, Path | None]:
    if args.source and args.input_file:
        raise WordHelperError("Positional argument and --input are mutually exclusive.")
    source_value = args.source or args.input_file

    if args.text is not None:
        if source_value:
            raise WordHelperError("--text cannot be combined with file input.")
        return _normalize_newlines(args.text), None

    if source_value is None or source_value == "-":
        if sys.stdin.isatty():
            raise WordHelperError(
                "Provide --text, a text file, or pipe input via standard input."
            )
        return _normalize_newlines(sys.stdin.read()), None

    source_path = Path(source_value).expanduser().resolve()
    if not source_path.is_file():
        raise WordHelperError(f"Input file not found: {source_path}")
    if source_path.suffix.lower() == ".docx":
        raise WordHelperError(
            "create expects a text file; use the read subcommand to read DOCX files."
        )
    return _read_text_file(source_path, args.encoding), source_path


def _default_output(source_path: Path | None) -> Path:
    if source_path:
        return source_path.with_suffix(".docx")
    return Path.cwd() / "output.docx"


def cmd_create(args) -> int:
    text, source_path = _read_input(args)
    if not text.strip():
        raise WordHelperError("Input text is empty.")

    resolved_format = resolve_format(args.format, source_path, text)
    output = (
        Path(args.output).expanduser()
        if args.output
        else _default_output(source_path)
    )
    if output.is_dir():
        raise WordHelperError(f"Output path is a directory: {output}")
    if not output.suffix:
        output = output.with_suffix(".docx")
    if output.suffix.lower() != ".docx":
        raise WordHelperError("Output file must have a .docx extension.")
    output = output.resolve()
    existed = output.exists()

    create_docx(
        text,
        output,
        title=args.title,
        author=args.author,
        date_text=args.date_text,
        fmt=resolved_format,
        source_path=source_path,
    )
    action = "Overwritten" if existed else "Generated"
    format_label = "Markdown" if resolved_format == "markdown" else "Plain text"
    print(f"{action}: {output}")
    print(f"Format: {format_label}")
    return 0


def cmd_read(args) -> int:
    path = Path(args.input).expanduser().resolve()
    data = extract_docx(
        path, include_all_parts=args.all
    )
    if args.json:
        content = json.dumps(data, ensure_ascii=False, indent=2)
    else:
        content = data["text"]

    if content and not content.endswith("\n"):
        content += "\n"

    if args.output:
        output = Path(args.output).expanduser().resolve()
        if output == path:
            raise WordHelperError("Output file cannot be the same as the input DOCX.")
        if output.is_dir():
            raise WordHelperError(f"Output path is a directory: {output}")
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(content, encoding="utf-8")
        print(f"Written: {output}", file=sys.stderr)
    else:
        sys.stdout.write(content)
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="word_helper.py",
        description=(
            "General-purpose Word tool: typeset DOCX from text or strings, "
            "or read all body text from a DOCX."
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Examples:\n"
            "  python word_helper.py create --text \"Hello, world\" -o hello.docx\n"
            "  python word_helper.py create --text 'Mass-energy equivalence $E=mc^2$' -o math.docx\n"
            "  python word_helper.py create --input notes.md -o notes.docx --title \"Project Notes\"\n"
            "  python word_helper.py read report.docx\n"
            "  python word_helper.py read report.docx --all --json\n"
        ),
    )
    subparsers = parser.add_subparsers(
        dest="command", required=True, metavar="{create,read}"
    )

    create_parser = subparsers.add_parser(
        "create",
        help="Generate a DOCX from a text file or a string",
        description="Typeset a string, text file, or standard input into DOCX (supports $...$ / $$...$$ math).",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    create_parser.add_argument(
        "source",
        nargs="?",
        help="Input text file; alternatively use --input, or pass - to read standard input",
    )
    create_parser.add_argument(
        "-i", "--input", dest="input_file", help="Path to an input text file"
    )
    create_parser.add_argument(
        "-t", "--text", help="String to typeset directly; mutually exclusive with file input"
    )
    create_parser.add_argument(
        "-o",
        "--output",
        help="Output .docx path; defaults to the input file name with .docx, or output.docx",
    )
    create_parser.add_argument("--title", help="Document title (shown at the top of the body)")
    create_parser.add_argument("--author", help="Author, shown below the title")
    create_parser.add_argument(
        "--date", dest="date_text", help="Date text, e.g. 2026-10-21"
    )
    create_parser.add_argument(
        "--format",
        choices=("auto", "plain", "markdown"),
        default="auto",
        help="Input format; default auto (.md files as Markdown, others detected from content)",
    )
    create_parser.add_argument(
        "--encoding",
        help="Text file encoding; by default tries UTF-8 / UTF-16 BOM / GB18030",
    )
    create_parser.set_defaults(func=cmd_create)

    read_parser = subparsers.add_parser(
        "read",
        help="Read the text content of a DOCX",
        description=(
            "Read DOCX body and tables in reading order; optionally read headers, footers, footnotes, and endnotes."
        ),
    )
    read_parser.add_argument("input", help="The .docx file to read")
    read_parser.add_argument(
        "--all",
        action="store_true",
        help="Also read headers, footers, first-page/even-page headers and footers, footnotes, and endnotes",
    )
    read_parser.add_argument(
        "--json",
        action="store_true",
        help="Output structured JSON instead of plain text",
    )
    read_parser.add_argument(
        "-o",
        "--output",
        help="Write the result to a UTF-8 text file instead of printing to the terminal",
    )
    read_parser.set_defaults(func=cmd_read)

    return parser


def _configure_console_encoding() -> None:
    """Fix garbled output when the console uses the UTF-8 code page but Python still writes GBK.

    Regular GBK consoles are left unchanged; no COM is involved and file content is not modified.
    """
    if os.name != "nt":
        return
    try:
        import ctypes

        if ctypes.windll.kernel32.GetConsoleOutputCP() != 65001:
            return
    except Exception:
        return

    for stream_name in ("stdin", "stdout", "stderr"):
        stream = getattr(sys, stream_name, None)
        if stream is None or not hasattr(stream, "reconfigure"):
            continue
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def main(argv: Sequence[str] | None = None) -> int:
    _configure_console_encoding()
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        return int(args.func(args))
    except WordHelperError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        print("Cancelled.", file=sys.stderr)
        return 130
    except BrokenPipeError:
        # Avoid flushing an already-closed pipe when the interpreter exits.
        try:
            null_fd = os.open(os.devnull, os.O_WRONLY)
            os.dup2(null_fd, sys.stdout.fileno())
            os.close(null_fd)
        except Exception:
            pass
        return 0
    except PermissionError as exc:
        print(f"Error: permission denied. {exc}", file=sys.stderr)
        return 2
    except Exception as exc:  # Last resort: avoid showing a long traceback to regular users.
        if os.environ.get("WORD_HELPER_DEBUG"):
            raise
        print(f"Error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
